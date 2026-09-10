import os
import json
import datetime
from typing import List, Optional

import cv2
from docx import Document
from docx.shared import Inches

from .frame_select import FrameChoice, choose_frame, choose_result_frame
from .redaction import RedactionConfig
from .utils import ensure_dir

JPEG_PARAMS = [int(cv2.IMWRITE_JPEG_QUALITY), 92]

TOAST_CAPTION = {
    "success": "Result (success message detected):",
    "error": "Result (error message detected):",
    "warning": "Result (warning message detected):",
}


def _write_capture(
    frame,
    path: str,
    redaction: Optional[RedactionConfig],
) -> str:
    if redaction is not None:
        frame = redaction.apply(frame)
    cv2.imwrite(path, frame, JPEG_PARAMS)
    return path


def _result_caption(choice: FrameChoice) -> str:
    if choice.has_toast:
        return TOAST_CAPTION.get(choice.toast.family, "Result (message detected):")
    return "Result:"


def build_evidence(
    video: str,
    markers: str,
    out_dir: str,
    title: str = "WHS Mobile – Test Evidence",
    skip_loading: bool = True,
    result_offsets: List[float] = None,
    redaction: Optional[RedactionConfig] = None,
    result_window: float = 2.5,
    detect_toast: bool = True,
):
    """
    Builds a Word doc from:
      - MP4 screen recording
      - step_markers.json with title/notes/is_loading

    For each marker:
      - capture 1 "action" screenshot around marker time
      - capture 1 "result" screenshot after the marker, preferring a frame that
        shows a result banner (toast) when `detect_toast` is on
      - apply the redaction rules, if any, before anything is written to disk
    """
    if result_offsets is None:
        result_offsets = [0.6, 1.2]

    ensure_dir(out_dir)

    with open(markers, "r", encoding="utf-8") as f:
        data = json.load(f)
    ms = data.get("markers", [])
    if not ms:
        raise RuntimeError("No markers found.")

    cap = cv2.VideoCapture(video)
    if not cap.isOpened():
        raise RuntimeError(f"Cannot open video: {video}")

    fps = cap.get(cv2.CAP_PROP_FPS) or 10.0

    run_stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    run_out = os.path.join(out_dir, f"evidence_{run_stamp}")
    ensure_dir(run_out)

    if redaction is not None and not redaction.is_empty():
        print(f"Redaction active: {redaction.describe()}")

    captures = []  # list of dicts per step

    step_no = 0
    for m in ms:
        if skip_loading and m.get("is_loading") is True:
            continue

        step_no += 1
        t = float(m.get("t", 0.0))
        step_title = (m.get("title") or f"Step {step_no}").strip()
        notes = (m.get("notes") or "").strip()

        # ACTION frame near marker
        action = choose_frame(cap, fps, t)
        action_path = None
        if action.frame is not None:
            action_path = _write_capture(
                action.frame,
                os.path.join(run_out, f"step_{step_no:02d}_action_{action.mode}.jpg"),
                redaction,
            )

        # RESULT frame AFTER marker (to catch the toast / success message)
        result = choose_result_frame(
            cap,
            fps,
            t,
            baseline=action.frame,
            offsets=result_offsets,
            window_sec=result_window,
            use_toast_detection=detect_toast,
        )

        result_path = None
        if result.frame is not None:
            name = f"step_{step_no:02d}_result_{result.mode}_plus{result.offset:.1f}s.jpg"
            result_path = _write_capture(result.frame, os.path.join(run_out, name), redaction)

        captures.append({
            "no": step_no,
            "t": t,
            "title": step_title,
            "notes": notes,
            "action_img": action_path,
            "result_img": result_path,
            "result_mode": result.mode,
            "result_offset": result.offset,
            "result_toast": result.toast.family if result.has_toast else None,
            "result_caption": _result_caption(result),
        })

    cap.release()

    # Build Word doc
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Video: {os.path.basename(video)}")
    doc.add_paragraph(f"Markers: {os.path.basename(markers)}")
    doc.add_paragraph("Per step: action screenshot + result screenshot (after marker) when available.")
    if redaction is not None and not redaction.is_empty():
        doc.add_paragraph(f"Redaction applied to every screenshot: {redaction.describe()}.")

    doc.add_heading("Steps", level=2)

    for c in captures:
        doc.add_heading(f"{c['no']}. {c['title']}", level=3)
        doc.add_paragraph(f"(Marker ~{c['t']:.1f}s)")
        if c["notes"]:
            doc.add_paragraph(c["notes"])

        if c["action_img"]:
            doc.add_paragraph("Action:")
            doc.add_picture(c["action_img"], width=Inches(6.3))

        if c["result_img"]:
            doc.add_paragraph(c["result_caption"])
            doc.add_picture(c["result_img"], width=Inches(6.3))

    out_doc = os.path.join(run_out, "WHS_Test_Evidence.docx")
    doc.save(out_doc)

    manifest = os.path.join(run_out, "steps.json")
    with open(manifest, "w", encoding="utf-8") as f:
        json.dump(
            {
                "video": os.path.basename(video),
                "markers": os.path.basename(markers),
                "generated": run_stamp,
                "redaction": redaction.describe() if redaction is not None else "none",
                "toast_detection": detect_toast,
                "steps": captures,
            },
            f,
            indent=2,
            ensure_ascii=False,
        )

    toasts = sum(1 for c in captures if c["result_toast"])
    print(f"Captured {len(captures)} step(s); result message detected on {toasts}.")
    print("Created:", out_doc)
    return out_doc
