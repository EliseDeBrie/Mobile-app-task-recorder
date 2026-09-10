"""Word output.

Two documents can be produced from the same recording:

* **task guide** - the replica of what D365 Task Recorder exports: the recording
  name, its description, then the numbered steps with their titles, notes and
  screenshots, grouped under the subtask headings.
* **evidence** - the test-evidence layout this tool started with, which pairs the
  action screenshot of each step with the screenshot of its result.
"""

import os
from dataclasses import dataclass
from typing import Dict, List, Optional

import cv2
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .instructions import PREFERRED
from .recording import INFO, STEP, SUBTASK_START, OutlineEntry, Recording

#: Handheld screenshots are portrait, so a full-width picture is unreadable.
PORTRAIT_WIDTH_IN = 2.9
LANDSCAPE_WIDTH_IN = 6.3
STEP_INDENT_IN = 0.3


@dataclass
class StepCapture:
    """The screenshots captured for one node of the recording."""

    node_index: int
    t: float = 0.0
    action_img: Optional[str] = None
    action_mode: str = "none"
    result_img: Optional[str] = None
    result_mode: str = "none"
    result_offset: float = 0.0
    result_toast: Optional[str] = None
    result_caption: str = "Result:"


def picture_width(path: str) -> Inches:
    """Size a picture so a portrait handheld capture stays legible on the page."""
    image = cv2.imread(path)
    if image is None:
        return Inches(LANDSCAPE_WIDTH_IN)
    height, width = image.shape[:2]
    return Inches(PORTRAIT_WIDTH_IN if height >= width else LANDSCAPE_WIDTH_IN)


def _add_picture(doc: Document, path: str, indent_in: float = 0.0) -> None:
    doc.add_picture(path, width=picture_width(path))
    paragraph = doc.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(indent_in)


def _add_annotation(doc: Document, text: str, indent_in: float, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(indent_in)
    run = paragraph.add_run(text)
    run.italic = italic


def _add_step_line(doc: Document, number: int, text: str, indent_in: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(indent_in)
    paragraph.paragraph_format.space_after = Pt(4)
    number_run = paragraph.add_run(f"{number}. ")
    number_run.bold = True
    paragraph.add_run(text)


def write_task_guide(
    recording: Recording,
    outline: List[OutlineEntry],
    captures: Dict[int, StepCapture],
    out_path: str,
    value_mode: str = PREFERRED,
    title: Optional[str] = None,
    include_result: bool = False,
) -> str:
    """Write the Task Recorder-style task guide."""
    doc = Document()
    doc.add_heading(title or recording.name, level=1)

    if recording.description:
        doc.add_paragraph(recording.description)

    for entry in outline:
        indent = STEP_INDENT_IN * entry.depth

        if entry.kind == SUBTASK_START:
            doc.add_heading(entry.node.name or "Subtask", level=min(2 + entry.depth, 4))
            continue

        if entry.node.title:
            _add_annotation(doc, entry.node.title, indent)

        _add_step_line(doc, entry.number, entry.node.instruction(value_mode), indent)

        capture = captures.get(entry.node_index) if entry.kind in (STEP, INFO) else None
        if capture and capture.action_img:
            _add_picture(doc, capture.action_img, indent)

        if include_result and capture and capture.result_img:
            _add_annotation(doc, capture.result_caption, indent)
            _add_picture(doc, capture.result_img, indent)

        if entry.node.note:
            _add_annotation(doc, entry.node.note, indent, italic=True)

    doc.save(out_path)
    return out_path


def write_evidence_document(
    recording: Recording,
    outline: List[OutlineEntry],
    captures: Dict[int, StepCapture],
    out_path: str,
    value_mode: str = PREFERRED,
    title: Optional[str] = None,
    video: str = "",
    markers: str = "",
    redaction_summary: str = "none",
) -> str:
    """Write the test-evidence document: action and result screenshot per step."""
    doc = Document()
    doc.add_heading(title or recording.name, level=1)

    if recording.description:
        doc.add_paragraph(recording.description)
    if video:
        doc.add_paragraph(f"Video: {os.path.basename(video)}")
    if markers:
        doc.add_paragraph(f"Recording: {os.path.basename(markers)}")
    doc.add_paragraph("Per step: action screenshot + result screenshot (after marker) when available.")
    if redaction_summary and redaction_summary != "none":
        doc.add_paragraph(f"Redaction applied to every screenshot: {redaction_summary}.")

    doc.add_heading("Steps", level=2)

    for entry in outline:
        if entry.kind == SUBTASK_START:
            doc.add_heading(entry.node.name or "Subtask", level=min(2 + entry.depth, 4))
            continue

        capture = captures.get(entry.node_index)
        doc.add_heading(f"{entry.number}. {entry.node.instruction(value_mode)}", level=3)

        if capture:
            doc.add_paragraph(f"(Marker ~{capture.t:.1f}s)")
        if entry.node.title:
            doc.add_paragraph(entry.node.title)
        if entry.node.note:
            doc.add_paragraph(entry.node.note)

        if capture and capture.action_img:
            doc.add_paragraph("Action:")
            _add_picture(doc, capture.action_img)
        if capture and capture.result_img:
            doc.add_paragraph(capture.result_caption)
            _add_picture(doc, capture.result_img)

    doc.save(out_path)
    return out_path
