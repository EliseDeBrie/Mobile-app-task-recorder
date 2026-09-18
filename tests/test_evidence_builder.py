import json
import os

import cv2
import pytest
from docx import Document

from conftest import add_banner, make_screen

from whs_recorder.evidence_builder import EVIDENCE, TASK_GUIDE, build_evidence
from whs_recorder.instructions import EXAMPLE
from whs_recorder.recording import InfoStep, Recording, Step, SubtaskEnd, SubtaskStart
from whs_recorder.redaction import RedactionConfig
from whs_recorder.utils import read_image
from whs_recorder.region import Region

GREEN = (0, 200, 0)


@pytest.fixture
def recording_path(tmp_path):
    r = Recording(name="Receive a purchase order line", description="WHS mobile receiving.")
    r.add(SubtaskStart(t=0.2, name="Open the work"))
    r.add(Step(t=1.5, action="scan", control="LP", value="LP000123", screen="Purchase receive",
               note="LP from the pallet label"))
    r.add(SubtaskEnd(t=2.0))
    r.add(Step(t=2.4, action="tap", control="OK", is_loading=True))
    r.add(Step(t=4.4, action="enter", control="Quantity", value="12"))
    r.add(InfoStep(t=5.4, text="Move the pallet to the staging lane"))
    path = tmp_path / "recording.json"
    r.save(str(path))
    return str(path)


@pytest.fixture
def video(tmp_path):
    """A 6-second recording with a success banner 0.9s after the first step."""
    path = tmp_path / "run.mp4"
    writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), 10.0, (360, 640))
    if not writer.isOpened():
        pytest.skip("No video writer available in this environment")

    for i in range(60):
        frame = make_screen(seed=i)
        if 24 <= i <= 28:
            frame = add_banner(frame, GREEN)
        writer.write(frame)
    writer.release()
    return str(path)


def read_run(out_dir):
    run_dir = os.path.join(out_dir, sorted(os.listdir(out_dir))[0])
    with open(os.path.join(run_dir, "steps.json"), encoding="utf-8") as f:
        return run_dir, json.load(f)


def shot(run_dir, manifest, index, key="action_img"):
    """The manifest names files; they live beside it."""
    return os.path.join(run_dir, manifest["steps"][index][key])


def doc_text(path):
    return [p.text for p in Document(path).paragraphs if p.text]


def test_task_guide_is_the_default_document(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    doc = build_evidence(video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True)

    assert os.path.basename(doc) == "Receive a purchase order line.docx"
    body = doc_text(doc)
    assert body[0] == "Receive a purchase order line"
    assert "Open the work" in body
    assert "1. In the LP field, scan 'LP000123'." in body
    assert "2. In the Quantity field, enter '12'." in body
    assert "3. Move the pallet to the staging lane." in body


def test_the_manifest_records_the_generated_instructions(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    build_evidence(video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True)

    _, manifest = read_run(str(out_dir))

    assert manifest["recording"] == "Receive a purchase order line"
    assert manifest["style"] == TASK_GUIDE
    assert manifest["value_mode"] == "preferred"
    assert [s["kind"] for s in manifest["steps"]] == ["subtask_start", "step", "step", "info"]
    assert manifest["steps"][1]["instruction"] == "In the LP field, scan 'LP000123'."
    assert manifest["steps"][1]["control"] == "LP"
    assert manifest["steps"][1]["screen"] == "Purchase receive"
    assert manifest["steps"][3].get("action_img") is None


def test_example_values_change_the_wording(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    doc = build_evidence(
        video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True, value_mode=EXAMPLE
    )

    assert "1. In the LP field, scan the value from the label." in doc_text(doc)


def test_a_step_screenshot_is_taken_for_every_step_but_not_for_info_steps(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    doc = build_evidence(video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True)

    assert len(Document(doc).inline_shapes) == 2


def test_loading_steps_are_skipped_on_request(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    build_evidence(video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=False)

    _, manifest = read_run(str(out_dir))
    assert [s["no"] for s in manifest["steps"] if s["kind"] == "step"] == [1, 2, 3]


def test_task_guide_can_include_the_result_screenshot(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    doc = build_evidence(
        video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True, include_result=True
    )

    body = doc_text(doc)
    assert "Result (success message detected):" in body
    assert len(Document(doc).inline_shapes) > 2


def test_evidence_style_still_pairs_action_and_result(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    doc = build_evidence(
        video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True, style=EVIDENCE
    )

    assert os.path.basename(doc) == "WHS_Test_Evidence.docx"
    body = doc_text(doc)
    assert "Action:" in body
    assert "1. In the LP field, scan 'LP000123'." in body

    run_dir, manifest = read_run(str(out_dir))
    step = manifest["steps"][1]
    assert step["result_toast"] == "success"
    assert step["result_mode"] == "toast"
    assert 0.8 <= step["result_offset"] <= 1.4
    saved = read_image(os.path.join(run_dir, step["result_img"]))
    assert (saved[560:600, :, 1] > 120).mean() > 0.8


def test_task_guide_skips_the_result_scan_it_does_not_need(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    build_evidence(video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True)

    run_dir, manifest = read_run(str(out_dir))
    assert not any("result" in name for name in os.listdir(run_dir))
    assert manifest["steps"][1]["result_img"] is None


def test_redaction_is_applied_to_every_saved_screenshot(video, recording_path, tmp_path):
    out_dir = tmp_path / "out"
    redaction = RedactionConfig.from_dict({"regions": [{"name": "user", "box": [0.0, 0.0, 1.0, 0.1]}]})

    build_evidence(
        video=video, markers=recording_path, out_dir=str(out_dir), skip_loading=True,
        redaction=redaction, style=EVIDENCE,
    )

    run_dir, manifest = read_run(str(out_dir))
    assert manifest["redaction"] == "1 region rule(s)"
    for key in ("action_img", "result_img"):
        assert (read_image(shot(run_dir, manifest, 1, key))[0:60] < 25).all()


def test_a_legacy_markers_file_still_builds(video, tmp_path):
    markers = tmp_path / "step_markers.json"
    markers.write_text(
        json.dumps({"markers": [{"t": 1.5, "title": "Scan the licence plate", "notes": "", "is_loading": False}]}),
        encoding="utf-8",
    )

    doc = build_evidence(video=video, markers=str(markers), out_dir=str(tmp_path / "out"))

    assert "1. Scan the licence plate." in doc_text(doc)


def test_an_empty_recording_is_rejected(video, tmp_path):
    empty = tmp_path / "empty.json"
    empty.write_text(json.dumps({"nodes": []}), encoding="utf-8")

    with pytest.raises(RuntimeError, match="No steps"):
        build_evidence(video=video, markers=str(empty), out_dir=str(tmp_path / "out"))


def test_unreadable_video_is_rejected(recording_path, tmp_path):
    with pytest.raises(RuntimeError, match="Cannot open video"):
        build_evidence(video=str(tmp_path / "nope.mp4"), markers=recording_path, out_dir=str(tmp_path / "out"))


def test_unknown_style_and_value_mode_are_rejected(video, recording_path, tmp_path):
    with pytest.raises(ValueError, match="Unknown style"):
        build_evidence(video=video, markers=recording_path, out_dir=str(tmp_path / "o"), style="fancy")

    with pytest.raises(ValueError, match="Unknown value mode"):
        build_evidence(video=video, markers=recording_path, out_dir=str(tmp_path / "o"), value_mode="loud")


@pytest.fixture
def recorded_screenshots(tmp_path):
    """A recording that carries its own screenshots, as `mark` now writes them."""
    shots = tmp_path / "receiving_screenshots"
    shots.mkdir()
    cv2.imwrite(str(shots / "step_01_action.png"), make_screen(seed=3))
    cv2.imwrite(str(shots / "step_01_result.png"), add_banner(make_screen(seed=4), GREEN))
    cv2.imwrite(str(shots / "step_02_action.png"), make_screen(seed=5))

    r = Recording(name="Receive a purchase order line", region=Region(100, 80, 360, 640, "select"))
    r.add(Step(t=1.5, action="scan", control="LP", value="LP000123",
               action_img="receiving_screenshots/step_01_action.png",
               result_img="receiving_screenshots/step_01_result.png",
               result_toast="success"))
    r.add(Step(t=3.0, action="tap", control="OK",
               action_img="receiving_screenshots/step_02_action.png"))
    path = tmp_path / "receiving.json"
    r.save(str(path))
    return str(path)


def test_a_recording_with_its_own_screenshots_needs_no_video(recorded_screenshots, tmp_path):
    doc = build_evidence(markers=recorded_screenshots, out_dir=str(tmp_path / "out"))

    assert "1. In the LP field, scan 'LP000123'." in doc_text(doc)
    assert len(Document(doc).inline_shapes) == 2

    _, manifest = read_run(str(tmp_path / "out"))
    assert manifest["video"] == ""
    assert manifest["region"] == "360x640 at 100,80 (select)"
    assert manifest["steps"][0]["action_img"].endswith("step_01_action.jpg")


def test_the_recorded_result_banner_keeps_its_caption(recorded_screenshots, tmp_path):
    doc = build_evidence(
        markers=recorded_screenshots, out_dir=str(tmp_path / "out"), include_result=True
    )

    assert "Result (success message detected):" in doc_text(doc)

    _, manifest = read_run(str(tmp_path / "out"))
    assert manifest["steps"][0]["result_toast"] == "success"
    assert manifest["steps"][0]["result_mode"] == "recorded"


def test_recorded_screenshots_are_redacted_at_build_time(recorded_screenshots, tmp_path):
    redaction = RedactionConfig.from_dict({"regions": [{"name": "user", "box": [0.0, 0.0, 1.0, 0.1]}]})

    build_evidence(markers=recorded_screenshots, out_dir=str(tmp_path / "out"), redaction=redaction)

    run_dir, manifest = read_run(str(tmp_path / "out"))
    assert (read_image(shot(run_dir, manifest, 0))[0:60] < 25).all()


def test_a_recording_with_neither_screenshots_nor_a_video_is_rejected(tmp_path):
    r = Recording(name="Empty handed")
    r.add(Step(t=1.0, action="tap", control="OK"))
    path = tmp_path / "recording.json"
    r.save(str(path))

    with pytest.raises(RuntimeError, match="--video"):
        build_evidence(markers=str(path), out_dir=str(tmp_path / "out"))


@pytest.fixture
def full_screen_video(tmp_path):
    """A screen recording of a whole 800x600 desktop, with the app in one corner."""
    path = tmp_path / "desktop.mp4"
    writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), 10.0, (800, 600))
    if not writer.isOpened():
        pytest.skip("No video writer available in this environment")
    for i in range(40):
        frame = make_screen(width=800, height=600, seed=i)
        writer.write(frame)
    writer.release()
    return str(path)


def test_video_frames_are_cropped_to_the_app_region(full_screen_video, tmp_path):
    r = Recording(name="Cropped", region=Region(100, 80, 200, 300, "select"))
    r.add(Step(t=1.5, action="tap", control="OK"))
    markers = tmp_path / "recording.json"
    r.save(str(markers))

    build_evidence(markers=str(markers), out_dir=str(tmp_path / "out"), video=full_screen_video)

    run_dir, manifest = read_run(str(tmp_path / "out"))
    assert read_image(shot(run_dir, manifest, 0)).shape[:2] == (300, 200)


def test_a_video_of_the_app_window_alone_is_left_uncropped(video, tmp_path):
    """The 360x640 video already shows only the app, so the region must not crop it again."""
    r = Recording(name="Not cropped", region=Region(0, 0, 360, 640, "select"))
    r.add(Step(t=1.5, action="tap", control="OK"))
    markers = tmp_path / "recording.json"
    r.save(str(markers))

    build_evidence(markers=str(markers), out_dir=str(tmp_path / "out"), video=video)

    run_dir, manifest = read_run(str(tmp_path / "out"))
    assert read_image(shot(run_dir, manifest, 0)).shape[:2] == (640, 360)


def test_the_document_can_also_be_placed_at_a_fixed_name(tmp_path, recorded_screenshots):
    """Every build has a run folder of its own; the launcher wants one file
    with a name that stays the same and is replaced each time."""
    wanted = tmp_path / "Receiving.docx"

    first = build_evidence(markers=recorded_screenshots, out_dir=str(tmp_path / "out"),
                           document=str(wanted))
    assert wanted.is_file()
    assert wanted.read_bytes() == open(first, "rb").read()

    before = wanted.stat().st_mtime_ns
    build_evidence(markers=recorded_screenshots, out_dir=str(tmp_path / "out"),
                   document=str(wanted))
    assert wanted.is_file()
    assert wanted.stat().st_mtime_ns >= before


def test_a_document_that_cannot_be_placed_says_where_the_build_is(tmp_path, recorded_screenshots, monkeypatch):
    """Word holding the previous copy open is the usual reason."""
    import shutil

    def locked(*_args, **_kwargs):
        raise PermissionError("being used by another process")

    monkeypatch.setattr(shutil, "copyfile", locked)

    with pytest.raises(RuntimeError, match="Close it if it is open"):
        build_evidence(markers=recorded_screenshots, out_dir=str(tmp_path / "out"),
                       document=str(tmp_path / "Receiving.docx"))
