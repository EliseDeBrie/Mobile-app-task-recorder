import cv2
import numpy as np
import pytest
from docx import Document
from docx.shared import Inches

from conftest import make_screen

from whs_recorder.instructions import EXAMPLE
from whs_recorder.recording import InfoStep, Recording, Step, SubtaskEnd, SubtaskStart
from whs_recorder.task_guide import (
    LANDSCAPE_WIDTH_IN,
    PORTRAIT_WIDTH_IN,
    StepCapture,
    picture_width,
    write_evidence_document,
    write_task_guide,
)


@pytest.fixture
def shot(tmp_path):
    path = tmp_path / "shot.jpg"
    cv2.imwrite(str(path), make_screen())  # 360x640, portrait
    return str(path)


@pytest.fixture
def recording():
    r = Recording(name="Receive a purchase order line", description="WHS mobile receiving.")
    r.add(SubtaskStart(t=0.2, name="Open the work"))
    r.add(Step(t=1.0, action="menu", control="Inbound"))
    r.add(SubtaskEnd(t=2.0))
    r.add(Step(t=3.0, action="scan", control="LP", value="LP000123",
               title="Check the pallet label first", note="A damaged label needs a reprint."))
    r.add(InfoStep(t=4.0, text="Move the pallet to the staging lane"))
    return r


def texts(doc):
    return [p.text for p in doc.paragraphs if p.text]


def styles(doc):
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text]


def test_task_guide_opens_with_the_recording_name_and_description(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out))

    doc = Document(str(out))
    assert styles(doc)[0] == ("Heading 1", "Receive a purchase order line")
    assert doc.paragraphs[1].text == "WHS mobile receiving."


def test_task_guide_renders_subtasks_as_headings(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out))

    assert ("Heading 2", "Open the work") in styles(Document(str(out)))


def test_task_guide_numbers_steps_and_generates_their_text(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out))

    body = texts(Document(str(out)))
    assert "1. Tap Inbound." in body
    assert "2. In the LP field, scan 'LP000123'." in body
    assert "3. Move the pallet to the staging lane." in body


def test_task_guide_honours_the_example_value_mode(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out), value_mode=EXAMPLE)

    body = texts(Document(str(out)))
    assert "2. In the LP field, scan the value from the label." in body


def test_title_sits_above_the_step_and_the_note_below_it(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out))

    body = texts(Document(str(out)))
    step = body.index("2. In the LP field, scan 'LP000123'.")
    assert body[step - 1] == "Check the pallet label first"
    assert body[step + 1] == "A damaged label needs a reprint."


def test_the_note_is_italic(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out))

    doc = Document(str(out))
    note = next(p for p in doc.paragraphs if p.text == "A damaged label needs a reprint.")
    assert note.runs[0].italic is True


def test_the_step_number_is_bold(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out))

    doc = Document(str(out))
    step = next(p for p in doc.paragraphs if p.text.startswith("1. "))
    assert step.runs[0].text == "1. "
    assert step.runs[0].bold is True


def test_the_title_parameter_overrides_the_recording_name(recording, tmp_path):
    out = tmp_path / "guide.docx"
    write_task_guide(recording, recording.outline(), {}, str(out), title="Receiving - UAT evidence")

    assert Document(str(out)).paragraphs[0].text == "Receiving - UAT evidence"


def test_screenshots_are_placed_under_their_step(recording, tmp_path, shot):
    outline = recording.outline()
    captures = {e.node_index: StepCapture(node_index=e.node_index, action_img=shot)
                for e in outline if e.kind == "step"}
    out = tmp_path / "guide.docx"
    write_task_guide(recording, outline, captures, str(out))

    assert len(Document(str(out)).inline_shapes) == 2


def test_the_result_screenshot_is_optional(recording, tmp_path, shot):
    outline = recording.outline()
    captures = {
        e.node_index: StepCapture(
            node_index=e.node_index, action_img=shot, result_img=shot,
            result_toast="success", result_caption="Result (success message detected):",
        )
        for e in outline if e.kind == "step"
    }

    without = tmp_path / "without.docx"
    write_task_guide(recording, outline, captures, str(without))
    assert len(Document(str(without)).inline_shapes) == 2

    with_result = tmp_path / "with.docx"
    write_task_guide(recording, outline, captures, str(with_result), include_result=True)
    doc = Document(str(with_result))
    assert len(doc.inline_shapes) == 4
    assert "Result (success message detected):" in texts(doc)


def test_an_info_step_renders_its_own_text_and_no_screenshot(recording, tmp_path, shot):
    """The builder captures nothing for an info step: no action was recorded."""
    outline = recording.outline()
    captures = {
        e.node_index: StepCapture(node_index=e.node_index, action_img=shot)
        for e in outline if e.kind == "step"
    }

    out = tmp_path / "guide.docx"
    write_task_guide(recording, outline, captures, str(out))

    doc = Document(str(out))
    assert "3. Move the pallet to the staging lane." in texts(doc)
    assert len(doc.inline_shapes) == 2


def test_portrait_screenshots_are_sized_down(shot, tmp_path):
    assert picture_width(shot) == Inches(PORTRAIT_WIDTH_IN)

    landscape = tmp_path / "wide.jpg"
    cv2.imwrite(str(landscape), np.zeros((400, 900, 3), dtype=np.uint8))
    assert picture_width(str(landscape)) == Inches(LANDSCAPE_WIDTH_IN)


def test_unreadable_image_falls_back_to_the_page_width(tmp_path):
    assert picture_width(str(tmp_path / "missing.jpg")) == Inches(LANDSCAPE_WIDTH_IN)


def test_evidence_document_keeps_its_provenance_and_pairs_the_screenshots(recording, tmp_path, shot):
    outline = recording.outline()
    captures = {
        e.node_index: StepCapture(
            node_index=e.node_index, t=e.node.t, action_img=shot, result_img=shot,
            result_caption="Result (success message detected):",
        )
        for e in outline if e.kind == "step"
    }

    out = tmp_path / "evidence.docx"
    write_evidence_document(
        recording, outline, captures, str(out),
        video="/runs/receiving.mp4", markers="/runs/recording.json",
        redaction_summary="1 region rule(s)",
    )

    doc = Document(str(out))
    body = texts(doc)
    assert "Video: receiving.mp4" in body
    assert "Recording: recording.json" in body
    assert "Redaction applied to every screenshot: 1 region rule(s)." in body
    assert ("Heading 3", "1. Tap Inbound.") in styles(doc)
    assert len(doc.inline_shapes) == 4
